import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[2] if Path(__file__).parent.name == 'docs' else Path(__file__).resolve().parents[2] / 'portafolio'

def make(name, pages, source):
    root = BASE / name
    if not root.exists(): return
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.6)
    sec.left_margin = sec.right_margin = Inches(.7)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.font.color.rgb = RGBColor.from_string('26332E')
    normal.paragraph_format.space_after = Pt(7)
    doc.styles['Caption'].font.name = 'Arial'
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.bold = False
    doc.styles['Caption'].font.color.rgb = RGBColor.from_string('58665F')
    for sty, size in [('Title', 27), ('Heading 1', 20), ('Heading 2', 12)]:
        doc.styles[sty].font.name = 'Arial'
        doc.styles[sty].font.size = Pt(size)
        doc.styles[sty].font.color.rgb = RGBColor.from_string('276749' if name == 'urban-demand' else '236C72')
    header = sec.header.paragraphs[0]
    header.text = name.upper().replace('-', ' ') + '  /  ANALYTICAL BRIEFING'
    header.style = 'Caption'
    footer = sec.footer.paragraphs[0]
    footer.text = 'Eliud Rojas Mendoza · Historical public-data study  |  '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for i, page in enumerate(pages):
        if i: doc.add_page_break()
        doc.add_paragraph(page['title'], 'Title' if i == 0 else 'Heading 1')
        for head, text in page['sections']:
            doc.add_paragraph(head, 'Heading 2')
            doc.add_paragraph(text)
        if page.get('figure'):
            doc.add_picture(str(root / 'reports/figures' / page['figure']), width=Inches(6.95))
            p = doc.add_paragraph(page['caption'], 'Caption'); p.paragraph_format.space_after = Pt(8)
        if page.get('table'):
            table = doc.add_table(rows=1, cols=len(page['table'][0])); table.style = 'Table Grid'
            for cell, txt in zip(table.rows[0].cells, page['table'][0]): cell.text = txt
            for row in page['table'][1:]:
                for cell, txt in zip(table.add_row().cells, row): cell.text = str(txt)
        doc.add_paragraph('Source: ' + source + ' · CC BY 4.0. Calculations: project R pipeline.', 'Caption')
    out = root / 'reports/downloads'; out.mkdir(exist_ok=True)
    doc.core_properties.title = name.replace('-', ' ').title() + ' — Analytical briefing'
    doc.core_properties.author = 'Eliud Rojas Mendoza'
    doc.core_properties.subject = 'Historical public-data analysis, methods and decisions'
    doc.core_properties.comments = ''
    doc.save(out / (name + '-briefing.docx'))
    print(name, '4-page briefing saved')

make('retail-insights', [
 {'title':'Retail Insights\nSales integrity and customer priorities', 'sections':[
 ('Business question', 'Which reporting issues and customer groups deserve attention before launching retention or product initiatives? The study audits all 541,909 Online Retail records, then connects a reconciled merchandise ledger with customer and product analysis.'),
 ('Measured scale', 'The primary ledger records £10,271,034.61 of merchandise purchases and £478,724.18 of negative-quantity credits or adjustments. Purchases less recorded credits are £9,792,310.43; this is not profit or verified accounting revenue.')], 'figure':'monthly_sales.png', 'caption':'Monthly merchandise purchase value. December 2011 covers only 1–9 December.'},
 {'title':'Investigate the ledger before acting', 'sections':[
 ('Explicit rules and reconciliation','Every source row receives one class. Administrative codes, invalid or nonpositive prices and cancellation inconsistencies remain visible. The complete row partition and aggregate amounts reconcile to the source.'),
 ('Sensitivity rather than silent deletion','The 5,268 exact repeated rows are preserved in primary results because no unique line identifier exists. Excluding later repeats reduces purchase value by £24,213.74. That difference is a reporting sensitivity, not recovered money.'),
 ('Credit candidates','582 groups contain exactly one purchase and one negative-quantity line sharing customer, code, date, absolute quantity and price, with £287,854.70 in candidate purchase value. A further 142 groups are ambiguous. These are investigation candidates, not proven return links.')], 'figure':'product_credits.png', 'caption':'High-credit merchandise codes. Credits alone do not establish defects, return reasons or stock losses.'},
 {'title':'Customers and observed cohorts', 'sections':[
 ('Concentration and coverage','The 4,334 identified buyers account for 85.3% of purchase value. The top 434 buyers contribute 61.3% of identified purchase value. Anonymous purchases are retained in the sales ledger but excluded from RFM.'),
 ('Retention priorities','RFM uses a fixed 10 December 2011 snapshot and consistent handling of ties. The segmentation identifies 384 at-risk buyers. Cohorts begin at the first purchase observed in this file, which may be later than a customer’s true first purchase.')], 'figure':'cohort_retention.png', 'caption':'Monthly active-buyer rates for first-observed cohorts. Only complete months enter this view; unobserved future cells remain blank.'},
 {'title':'A practical decision sequence', 'sections':[
 ('1. Resolve reporting questions','Review high-value credit candidates and administrative codes with invoice owners. Require an invoice-line identifier and a product master before adopting automatic matching or revised revenue rules.'),
 ('2. Design a measurable retention pilot','Review contact eligibility and consent, randomise comparable eligible buyers into contact and holdout groups, and measure incremental contribution after campaign costs. The scenario calculator uses editable assumptions and does not claim realised benefits.'),
 ('3. Make the analysis repeatable','Restore renv, verify source checksums, execute rule fixtures and run the complete R pipeline. Use the full report and aggregate workbook for traceability; individual customer identifiers are excluded from the public explorer.')],
 'table':[['Deliverable','Purpose'],['10 figures + complete HTML','Methods, reconciliations and interpretation'],['Interactive white explorer','Markets, month ranges, product ranking and sensitivity'],['Excel aggregate workbook','Auditable tables and an editable pilot scenario'],['RStudio project + source records','Repeatable transformations and model rules']]}
], 'Chen, D. (2015), Online Retail, UCI. https://doi.org/10.24432/C5BW33')
make('urban-demand', [
 {'title':'Urban Demand\nBicycle rentals and operational planning', 'sections':[
 ('Business question','How does system-wide rental demand vary by calendar, hour and weather, and how accurately can an interpretable model estimate later recorded counts? The complete archive contains 17,379 hourly observations and 731 daily observations from 2011–2012.'),
 ('Measured scale','The source records 3,292,679 rentals, with registered riders accounting for 81.2%. Working-day and weekend profiles differ, providing evidence for workload discussions when combined with service times and capacity.')], 'figure':'hourly_profiles.png', 'caption':'Mean recorded rentals per observed hour, separated by day type. These are system-wide patterns, not station inventory recommendations.'},
 {'title':'Select models using earlier data', 'sections':[
 ('Fair temporal comparison','Four candidates are compared using three expanding training windows and later validation quarters. Average validation MAE selects the spline negative-binomial model. The final October–December 2012 test does not participate in selection.'),
 ('Control information leakage','Calendar, trend and observed weather enter the model. Casual and registered counts are excluded because they are target components. Spline transformations are fitted within each training window.')], 'figure':'model_validation.png', 'caption':'Average MAE over three validation folds; lower is better. The reference uses historical means for hour and day type.'},
 {'title':'Final-test performance and uncertainty', 'sections':[
 ('Out-of-time estimates','On 2,168 final-test hours, the selected model achieves MAE 47.05 rentals/hour versus 79.95 for the reference, a 41.2% reduction. WAPE is 21.5%. Evaluation uses contemporaneous observed weather and does not establish advance-forecast accuracy.'),
 ('Intervals need calibration','Conditional nominal 95% count intervals cover only 87.6% of test outcomes. Their width omits parameter and weather-forecast uncertainty. Hourly and weather-specific diagnostics expose performance differences hidden by averages.')], 'figure':'heldout_daily.png', 'caption':'Daily sums of actual counts and saved hourly estimates on the final quarter. These are conditional estimates evaluated on later dates.'},
 {'title':'From evidence to a planning pilot', 'sections':[
 ('1. Preserve missingness','There are 165 absent hour records across 76 days. Observed hourly totals reconcile exactly with every daily total, but the cause of the missing records is unknown. Do not convert absent hours into verified zero demand.'),
 ('2. Evaluate a real forecasting horizon','Use archived weather forecasts available at the decision time, compare against the same reference on later observations, and calibrate uncertainty before choosing buffers. A seeded seven-day block bootstrap quantifies paired error differences for this historical test.'),
 ('3. Add operational data','Station movements, capacities, stock-out duration, route costs, staffing service times and costs are required to evaluate deployment decisions. System totals do not measure denied demand, savings or where bicycles should be placed.')],
 'table':[['Deliverable','Purpose'],['10 figures + complete HTML','Patterns, temporal validation and diagnostic limits'],['Interactive white explorer','Year, day type, weather and final-test month'],['Excel aggregate workbook','Validation, saved estimates and source coverage'],['RStudio project + locked packages','Source verification and repeatable modelling']]}
], 'Fanaee-T, H. (2013), Bike Sharing, UCI. https://doi.org/10.24432/C5W894')
